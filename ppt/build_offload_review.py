#!/usr/bin/env python3
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT_DIR = Path(__file__).resolve().parent
OUT_FILE = OUT_DIR / "llama_cpp_moe_offload_experiment_review_zh.pptx"

SW = 13.333
SH = 7.5

BG = "F7F8F5"
PAPER = "FFFFFF"
INK = "1D252C"
MUTED = "68737D"
GRID = "DDE2E3"
TEAL = "0F8B8D"
TEAL_LIGHT = "D9EEEE"
BLUE = "3D7EA6"
BLUE_LIGHT = "DFEAF2"
GREEN = "2E8B57"
GREEN_LIGHT = "DDEEE4"
AMBER = "F4B942"
AMBER_LIGHT = "FCEEC6"
RED = "D1495B"
RED_LIGHT = "F6DEE2"
VIOLET = "7B5EA7"
VIOLET_LIGHT = "EAE4F2"
CHARCOAL = "26343D"

FONT_CN = "Microsoft YaHei"
FONT_MONO = "Aptos Mono"


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def set_bg(slide, color=BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def add_shape(slide, kind, x, y, w, h, fill=PAPER, line=GRID, radius=False):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else kind
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(1)
    return shape


def add_text(
        slide, text, x, y, w, h, size=18, color=INK, bold=False,
        font=FONT_CN, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
        margin=0.05, fit=False, line_spacing=1.0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    tf.word_wrap = True
    if fit:
        tf.fit_text(max_size=size)
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.line_spacing = line_spacing
    run = p.runs[0]
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def add_rich_text(slide, runs, x, y, w, h, size=18, color=INK, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = valign
    tf.word_wrap = True
    p = tf.paragraphs[0]
    for item in runs:
        run = p.add_run()
        run.text = item[0]
        run.font.name = item[4] if len(item) > 4 else FONT_CN
        run.font.size = Pt(item[1] if len(item) > 1 else size)
        run.font.bold = item[2] if len(item) > 2 else False
        run.font.color.rgb = rgb(item[3] if len(item) > 3 else color)
    return box


def add_bullets(slide, items, x, y, w, h, size=16, color=INK, bullet_color=None, spacing=6):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(0.03)
    tf.margin_right = Inches(0.03)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = FONT_CN
        p.font.size = Pt(size)
        p.font.color.rgb = rgb(color)
        p.space_after = Pt(spacing)
        p.line_spacing = 1.06
        p.text = "•  " + item
        if p.runs:
            p.runs[0].font.color.rgb = rgb(bullet_color or color)
    return box


def add_pill(slide, text, x, y, w, color=TEAL, fill=TEAL_LIGHT, size=11):
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, 0.34, fill, fill, True)
    add_text(slide, text, x + 0.04, y + 0.02, w - 0.08, 0.28, size, color, True,
             align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)


def add_title(slide, title, kicker=None, section_color=TEAL):
    if kicker:
        add_text(slide, kicker.upper(), 0.72, 0.34, 3.7, 0.28, 10, section_color, True, margin=0)
    add_text(slide, title, 0.72, 0.68 if kicker else 0.46, 11.9, 0.62, 27, INK, True, margin=0)
    y = 1.33 if kicker else 1.13
    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.72), Inches(y), Inches(0.62), Inches(0.04))
    line.fill.solid()
    line.fill.fore_color.rgb = rgb(section_color)
    line.line.fill.background()
    return y


def add_footer(slide, source, page):
    y = 7.18
    add_text(slide, source, 0.72, y, 10.9, 0.18, 8.5, MUTED, margin=0)
    add_text(slide, f"{page:02d}", 12.05, y - 0.01, 0.55, 0.2, 9, MUTED, True,
             align=PP_ALIGN.RIGHT, margin=0)


def add_card(slide, x, y, w, h, title=None, fill=PAPER, line=GRID, accent=None):
    card = add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, w, h, fill, line, True)
    if accent:
        stripe = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(0.07), Inches(h))
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = rgb(accent)
        stripe.line.fill.background()
    if title:
        add_text(slide, title, x + 0.23, y + 0.18, w - 0.42, 0.35, 15, INK, True, margin=0)
    return card


def add_metric(slide, value, label, x, y, w, color=TEAL, note=None):
    add_text(slide, value, x, y, w, 0.55, 28, color, True, margin=0)
    add_text(slide, label, x, y + 0.58, w, 0.34, 12, INK, True, margin=0)
    if note:
        add_text(slide, note, x, y + 0.92, w, 0.35, 10, MUTED, margin=0)


def add_arrow(slide, x1, y1, x2, y2, color=MUTED, width=2.0):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(width)
    line.line.end_arrowhead = True
    return line


def add_bar_chart(slide, categories, series, x, y, w, h, max_value=None, label_size=10, value_fmt="{:.2f}"):
    # series: [(name, values, color)]
    left = x + 0.72
    top = y + 0.30
    plot_w = w - 0.92
    plot_h = h - 0.86
    all_values = [v for _, values, _ in series for v in values]
    max_v = max_value or max(all_values) * 1.12
    n_cat = len(categories)
    group_w = plot_w / n_cat
    bar_gap = 0.05
    bar_w = min(0.34, (group_w - 0.14) / max(1, len(series)))

    axis = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(left), Inches(top + plot_h),
                                      Inches(left + plot_w), Inches(top + plot_h))
    axis.line.color.rgb = rgb(GRID)
    axis.line.width = Pt(1)
    for ci, cat in enumerate(categories):
        group_x = left + ci * group_w + group_w / 2
        total_w = len(series) * bar_w + (len(series) - 1) * bar_gap
        start_x = group_x - total_w / 2
        for si, (_, values, color) in enumerate(series):
            val = values[ci]
            bh = plot_h * val / max_v
            bx = start_x + si * (bar_w + bar_gap)
            by = top + plot_h - bh
            rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(bx), Inches(by), Inches(bar_w), Inches(bh))
            rect.fill.solid()
            rect.fill.fore_color.rgb = rgb(color)
            rect.line.fill.background()
            add_text(slide, value_fmt.format(val), bx - 0.15, by - 0.23, bar_w + 0.30, 0.2,
                     label_size - 1, color, True, align=PP_ALIGN.CENTER, margin=0)
        add_text(slide, str(cat), group_x - group_w / 2, top + plot_h + 0.08, group_w, 0.25,
                 label_size, INK, True, align=PP_ALIGN.CENTER, margin=0)
    legend_x = x + 0.72
    for name, _, color in series:
        dot = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(legend_x), Inches(y), Inches(0.12), Inches(0.12))
        dot.fill.solid()
        dot.fill.fore_color.rgb = rgb(color)
        dot.line.fill.background()
        add_text(slide, name, legend_x + 0.17, y - 0.04, 1.7, 0.22, 9, MUTED, margin=0)
        legend_x += 1.65


def add_hbar_chart(slide, labels, values, x, y, w, h, color=TEAL, max_value=None, suffix="", baseline=None):
    max_v = max_value or max(values) * 1.08
    row_h = h / len(labels)
    label_w = 1.35
    value_w = 0.85
    plot_w = w - label_w - value_w
    for i, (label, value) in enumerate(zip(labels, values)):
        cy = y + i * row_h
        add_text(slide, label, x, cy + 0.07, label_w - 0.1, row_h - 0.08, 11, INK, True,
                 valign=MSO_ANCHOR.MIDDLE, margin=0)
        bg = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x + label_w), Inches(cy + 0.17),
                                    Inches(plot_w), Inches(0.22))
        bg.fill.solid()
        bg.fill.fore_color.rgb = rgb("E9EDEE")
        bg.line.fill.background()
        bw = max(0.02, plot_w * value / max_v)
        bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x + label_w), Inches(cy + 0.17),
                                     Inches(bw), Inches(0.22))
        bar.fill.solid()
        bar.fill.fore_color.rgb = rgb(color if baseline is None or value <= baseline else RED)
        bar.line.fill.background()
        add_text(slide, f"{value:g}{suffix}", x + label_w + plot_w + 0.08, cy + 0.05,
                 value_w - 0.05, row_h - 0.05, 11, color, True, valign=MSO_ANCHOR.MIDDLE, margin=0)


def add_table(slide, headers, rows, x, y, w, h, widths=None, header_fill=CHARCOAL, highlight_rows=None):
    ncols = len(headers)
    nrows = len(rows) + 1
    widths = widths or [w / ncols] * ncols
    scale = w / sum(widths)
    widths = [cw * scale for cw in widths]
    rh = h / nrows
    for r in range(nrows):
        cx = x
        for c in range(ncols):
            cw = widths[c]
            if r == 0:
                fill = header_fill
                text_color = PAPER
                bold = True
                value = headers[c]
            else:
                is_hi = highlight_rows and (r - 1) in highlight_rows
                fill = AMBER_LIGHT if is_hi else (PAPER if r % 2 else "F0F3F2")
                text_color = INK
                bold = bool(is_hi) or c == 0
                value = rows[r - 1][c]
            cell = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(cx), Inches(y + r * rh),
                                          Inches(cw), Inches(rh))
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(fill)
            cell.line.color.rgb = rgb(BG if r == 0 else GRID)
            cell.line.width = Pt(0.6)
            add_text(slide, str(value), cx + 0.05, y + r * rh + 0.02, cw - 0.1, rh - 0.04,
                     10.5 if r else 10.2, text_color, bold,
                     align=PP_ALIGN.CENTER if c > 0 else PP_ALIGN.LEFT,
                     valign=MSO_ANCHOR.MIDDLE, margin=0)
            cx += cw


def new_slide(prs, bg=BG):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, bg)
    return slide


def make_deck():
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    prs.core_properties.title = "llama.cpp MoE Offload 实验实现与分支尝试复盘"
    prs.core_properties.subject = "共享 scratch、动态专家缓存、准入策略、投机解码、MTP 与 DRAM 流式传输"
    prs.core_properties.author = "Codex（基于仓库代码与实验产物整理）"
    prs.core_properties.keywords = "llama.cpp, MoE, offload, ubatch, EAMC, speculative decoding, MTP"
    page = 0

    def finish(slide, source):
        nonlocal page
        page += 1
        add_footer(slide, source, page)

    # 1. Cover
    slide = new_slide(prs, CHARCOAL)
    add_pill(slide, "CODE + BRANCHES + BENCHMARK ARTIFACTS", 0.78, 0.62, 3.32, AMBER, "514933", 10)
    add_text(slide, "llama.cpp MoE Offload", 0.78, 1.30, 11.6, 0.75, 31, PAPER, True, margin=0)
    add_text(slide, "实验实现与分支尝试复盘", 0.78, 2.10, 11.6, 0.72, 29, AMBER, True, margin=0)
    add_text(slide, "从共享 scratch、动态专家缓存，到准入策略、投机解码、MTP 与 DRAM 流式传输",
             0.80, 3.08, 10.7, 0.66, 17, "DDE5E8", margin=0)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0.80, 4.28, 11.72, 0.02, TEAL, TEAL)
    add_metric(slide, "5", "实验分支", 0.82, 4.66, 1.7, AMBER, "0714 / 0716 / 0719 / 0720 / 0724")
    add_metric(slide, "12", "关键提交", 3.05, 4.66, 1.7, TEAL, "2026-07-16 至 2026-07-27")
    add_metric(slide, "40 / 41", "MoE 层", 5.35, 4.66, 2.2, BLUE, "Q4 基线 / MTP 变体")
    add_metric(slide, "256", "专家 / 层", 8.10, 4.66, 2.0, GREEN, "top-k = 8")
    add_text(slide, "代码基线：upstream/master@cb489bc0f  ·  复盘日期：2026-07-27",
             0.80, 6.68, 10.8, 0.25, 10, "AAB7BD", margin=0)
    finish(slide, "仓库：/home/chaoyang.zhang/docker_workspace/projects/llama.cpp.clean")

    # 2. Executive summary
    slide = new_slide(prs)
    add_title(slide, "先给结论：基础架构成立，局部策略有效，跨 token 重叠尚未成立", "EXECUTIVE SUMMARY")
    items = [
        ("可保留", GREEN, GREEN_LIGHT, "共享 scratch + 全专家活动轴", "让每层任意 256 个专家都可被当前 ubatch 消费；6 GiB 缓存下仍能把 ubatch 放到 1024。"),
        ("已证实", TEAL, TEAL_LIGHT, "大 ubatch + 频次准入", "ubatch 128→1024 使 6 GiB TTFT/token 降 54.8%；频次准入使前 8 个 decode token 的 miss 降 37.6%。"),
        ("需重做", RED, RED_LIGHT, "n-gram 投机与 gate-only 流水", "候选 token 放大专家 I/O；DRAM gate-only 流式在 64–1024 token 全部退化，P0 清理控制开销后仍退化。"),
        ("待验证", VIOLET, VIOLET_LIGHT, "MTP offload", "功能路径、MMQ 边界修复和分阶段指标已完成；但缺少同量化、同 prompt、同二进制的 no-MTP A/B。"),
    ]
    for i, (tag, color, fill, title, body) in enumerate(items):
        x = 0.78 + (i % 2) * 6.12
        y = 1.62 + (i // 2) * 2.33
        add_card(slide, x, y, 5.72, 1.90, fill=PAPER, accent=color)
        add_pill(slide, tag, x + 0.25, y + 0.22, 0.88, color, fill, 10)
        add_text(slide, title, x + 1.32, y + 0.20, 4.05, 0.38, 16, INK, True, margin=0)
        add_text(slide, body, x + 0.25, y + 0.78, 5.05, 0.76, 13, MUTED, margin=0)
    add_text(slide, "最重要的工程判断", 0.80, 6.40, 1.65, 0.28, 11, RED, True, margin=0)
    add_text(slide, "下一轮不应继续微调 gate-only 分组；应先建立 token 等价性门槛，再做完整 gate/up→activation→down 的 ready-set executor。",
             2.48, 6.36, 9.95, 0.43, 13, INK, True, margin=0)
    finish(slide, "综合：539f4cbf4 / 8b4de7c0e / 94fdf126b / fde549896 / a03c4fa82")

    # 3. Scope
    slide = new_slide(prs)
    add_title(slide, "复盘口径：代码、提交和原始结果三方互证", "SCOPE & EVIDENCE", BLUE)
    add_card(slide, 0.78, 1.54, 3.74, 4.95, title="阅读范围", accent=BLUE)
    add_bullets(slide, [
        "5 条本地实验分支及其共同祖先",
        "MoE offload 运行时、CUDA I/O、MMQ、repack/bench 工具",
        "summary / CSV / log / RESULTS.md 与复现实验脚本",
        "按提交父子关系区分共享实现与分支独有尝试",
    ], 1.03, 2.10, 3.18, 2.72, 13.2, MUTED, BLUE, 9)
    add_pill(slide, "证据等级 A", 1.04, 5.25, 1.28, GREEN, GREEN_LIGHT, 10)
    add_text(slide, "同二进制开关或交替 A/B", 2.46, 5.27, 1.70, 0.27, 10.5, INK, True, margin=0)
    add_pill(slide, "B", 1.04, 5.71, 0.47, TEAL, TEAL_LIGHT, 10)
    add_text(slide, "单点/矩阵，配置清楚", 1.66, 5.73, 2.25, 0.27, 10.5, INK, True, margin=0)
    add_pill(slide, "C", 1.04, 6.13, 0.47, AMBER, AMBER_LIGHT, 10)
    add_text(slide, "探索性，缺少严格对照", 1.66, 6.15, 2.35, 0.27, 10.5, INK, True, margin=0)

    add_card(slide, 4.78, 1.54, 3.65, 4.95, title="Q4 主实验平台", accent=TEAL)
    add_metric(slide, "NVIDIA L20", "45,457 MiB", 5.04, 2.12, 2.65, TEAL)
    add_metric(slide, "20.60 GiB", "Qwen3.6-35B-A3B Q4_K", 5.04, 3.38, 2.65, BLUE)
    add_metric(slide, "40 × 256 × 3", "MoE 层 × 专家 × 权重种类", 5.04, 4.67, 2.90, GREEN)
    add_text(slide, "主口径：pp1024 / tg1024 / greedy / cache reset；6 GiB 与 8 GiB 专家缓存。",
             5.04, 5.95, 2.95, 0.48, 11.5, MUTED, margin=0)

    add_card(slide, 8.69, 1.54, 3.86, 4.95, title="独立标注的例外", accent=VIOLET)
    add_pill(slide, "MTP", 8.96, 2.12, 0.68, VIOLET, VIOLET_LIGHT, 10)
    add_text(slide, "Q8_0 · 41 个 MoE 层 · blob max 1.06 MiB", 9.83, 2.13, 2.30, 0.50, 12, INK, True, margin=0)
    add_pill(slide, "上限", 8.96, 3.09, 0.68, AMBER, AMBER_LIGHT, 10)
    add_text(slide, "llama-bench 全驻留结果仅作性能天花板，不与 moe-bench 视为严格 A/B。", 9.83, 3.08, 2.30, 0.82, 11.5, MUTED, margin=0)
    add_pill(slide, "正确性", 8.96, 4.26, 0.88, RED, RED_LIGHT, 10)
    add_text(slide, "模式内 repeat 一致 ≠ baseline/spec token 等价；两者必须分别检查。", 9.99, 4.24, 2.15, 0.82, 11.5, MUTED, margin=0)
    add_pill(slide, "时间", 8.96, 5.45, 0.68, BLUE, BLUE_LIGHT, 10)
    add_text(slide, "提交时间为 2026-07-16 至 07-27；结果按仓库内记录原样引用。", 9.83, 5.43, 2.30, 0.80, 11.5, MUTED, margin=0)
    finish(slide, "AGENTS.md；git log --all；models/*.moe.gguf.json；benchmark-results/**")

    # 4. Branch topology
    slide = new_slide(prs)
    add_title(slide, "分支演进不是一条直线：三条优化路线从共同核心分叉", "BRANCH TOPOLOGY", VIOLET)
    y0 = 2.14
    nodes = [
        (0.85, "master", "cb489bc0f", CHARCOAL),
        (2.40, "共享框架", "539f4cbf4", TEAL),
        (4.14, "频次准入", "8b4de7c0e", GREEN),
        (5.82, "驱逐微调", "c09907bec", GREEN),
        (7.42, "4K 测试", "3821d9835", BLUE),
    ]
    for i in range(len(nodes) - 1):
        add_arrow(slide, nodes[i][0] + 1.05, y0 + 0.36, nodes[i + 1][0] - 0.08, y0 + 0.36, MUTED, 1.5)
    for x, label, commit, color in nodes:
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y0, 1.15, 0.74, PAPER, color, True)
        add_text(slide, label, x + 0.05, y0 + 0.10, 1.05, 0.24, 10.5, INK, True,
                 align=PP_ALIGN.CENTER, margin=0)
        add_text(slide, commit, x + 0.05, y0 + 0.41, 1.05, 0.18, 8.3, color, True,
                 font=FONT_MONO, align=PP_ALIGN.CENTER, margin=0)

    branches = [
        (3.00, 3.60, 2.00, "0719_speculative_decoding", "94fdf126b · n-gram 投机", RED),
        (5.30, 4.55, 2.22, "0720_mtp", "0edf6febc → 941f360ef → fde549896", VIOLET),
        (7.95, 3.60, 1.72, "0716_eamc", "4306ebc28 → 5dd375005", BLUE),
        (9.95, 4.55, 2.50, "0724_transfer", "2dcac38ea → a03c4fa82", AMBER),
    ]
    # Branch stems.
    add_arrow(slide, 6.25, y0 + 0.75, 4.05, 3.56, RED, 1.7)
    add_arrow(slide, 7.99, y0 + 0.75, 6.40, 4.50, VIOLET, 1.7)
    add_arrow(slide, 7.99, y0 + 0.75, 8.80, 3.56, BLUE, 1.7)
    add_arrow(slide, 8.85, 3.99, 11.00, 4.50, AMBER, 1.7)
    for x, y, w, name, desc, color in branches:
        add_card(slide, x, y, w, 1.15, fill=PAPER, accent=color)
        add_text(slide, name, x + 0.20, y + 0.17, w - 0.32, 0.30, 11.5, color, True, margin=0)
        add_text(slide, desc, x + 0.20, y + 0.58, w - 0.32, 0.35, 9.3, MUTED, margin=0)
    add_pill(slide, "0714_ubatch_new", 1.02, 3.95, 1.52, TEAL, TEAL_LIGHT, 10)
    add_text(slide, "仅包含共同框架提交 539f4cbf4", 1.02, 4.40, 1.85, 0.58, 10.5, MUTED, margin=0)
    add_text(slide, "阅读原则", 0.85, 6.27, 0.82, 0.26, 10.5, INK, True, margin=0)
    add_text(slide, "0719 与 0720 是互斥的投机路线；0724 继承 0716 的测量与准入改动，不继承 0719/0720。",
             1.82, 6.22, 10.55, 0.40, 12.5, INK, margin=0)
    finish(slide, "git merge-base / git log：所有实验分支共同基线 master@cb489bc0f")

    # 5. Problem definition
    slide = new_slide(prs)
    add_title(slide, "目标不是“让模型能跑”，而是在小显存预算下控制专家换入成本", "PROBLEM MODEL", RED)
    add_card(slide, 0.78, 1.52, 3.52, 4.98, title="原始约束", accent=RED)
    add_metric(slide, "20.60 GiB", "Q4_K 模型文件", 1.04, 2.14, 2.55, RED)
    add_metric(slide, "256", "每层路由专家总数", 1.04, 3.40, 2.20, VIOLET)
    add_metric(slide, "8", "每 token 实际激活专家", 1.04, 4.66, 2.20, TEAL)
    add_text(slide, "专家权重占大头，但每个 token 只访问少量专家；这给“持久缓存 + 按需换入”留下空间。",
             1.04, 5.83, 2.86, 0.58, 11.7, MUTED, margin=0)

    add_card(slide, 4.58, 1.52, 3.76, 4.98, title="实验目标", accent=TEAL)
    add_bullets(slide, [
        "把专家权重从常规 loader 中截获",
        "用固定 VRAM budget 保存热点专家",
        "miss 时从 SSD/DRAM 读取并异步 H2D",
        "保持上游 ggml_mul_mat_id 单一 3D 权重接口",
        "量化 TTFT / TPOT / hit / bytes / stall",
    ], 4.87, 2.16, 3.12, 2.75, 13.2, MUTED, TEAL, 8)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 4.88, 5.30, 3.08, 0.78, TEAL_LIGHT, TEAL, True)
    add_text(slide, "核心设计取舍", 5.08, 5.44, 0.95, 0.22, 10, TEAL, True, margin=0)
    add_text(slide, "不改双源 CUDA kernel", 6.14, 5.39, 1.52, 0.32, 13.2, INK, True, margin=0)

    add_card(slide, 8.62, 1.52, 3.92, 4.98, title="观察到的交换", accent=AMBER)
    add_metric(slide, "9.74 GB", "6 GiB 专家缓存时进程峰值估算", 8.92, 2.15, 3.05, AMBER)
    add_metric(slide, "26.90 ms", "基线 offload TPOT", 8.92, 3.52, 2.70, RED)
    add_metric(slide, "158.95 tok/s", "全驻留 llama-bench decode 天花板", 8.92, 4.84, 3.02, GREEN)
    add_text(slide, "结论：容量目标已达到，但性能主要受 callback、读盘/H2D 和小粒度同步支配。",
             8.92, 5.94, 3.06, 0.46, 11.5, INK, True, margin=0)
    finish(slide, "models/Qwen3.6-35B-A3B-UD-Q4_K_M.moe.gguf.json；new-result/non-offload-single-gpu")

    # 6. Repack contract
    slide = new_slide(prs)
    add_title(slide, "离线重排建立了运行时可随机访问的专家索引", "BASE IMPLEMENTATION · REPACK", TEAL)
    stages = [
        (0.82, "原始 GGUF", "融合 expert tensor\n[dim0, dim1, 256]", CHARCOAL),
        (3.42, "llama-moe-repack", "扫描 MoE tensor\n计算每专家 byte range", TEAL),
        (6.45, "v2 元数据", "layer_ids / n_experts\nblob.table[offset,size]", BLUE),
        (9.38, "运行时 loader", "inspect manifest\n按 (L,E,kind) 定位", GREEN),
    ]
    for i, (x, title, body, color) in enumerate(stages):
        add_card(slide, x, 2.02, 2.06, 1.72, fill=PAPER, accent=color)
        add_text(slide, title, x + 0.22, 2.28, 1.62, 0.32, 14, color, True, margin=0)
        add_text(slide, body, x + 0.22, 2.78, 1.62, 0.58, 11.2, MUTED, margin=0)
        if i < len(stages) - 1:
            add_arrow(slide, x + 2.16, 2.88, stages[i + 1][0] - 0.10, 2.88, MUTED, 1.6)
    add_card(slide, 0.82, 4.34, 7.34, 1.72, title="运行时契约", accent=TEAL)
    add_bullets(slide, [
        "元数据：moe_offload.version=2，40 层，256 experts/layer，max blob=860,160 B",
        "原专家 tensor 被 mark_tensor_unloaded；loader 创建无文件 backing 的 GPU slot tensor",
        "同一 kind 在不同层可能是不同量化类型，因此按 (kind, dtype) 分配 global tensor",
    ], 1.08, 4.88, 6.68, 0.98, 11.5, MUTED, TEAL, 4)
    add_card(slide, 8.47, 4.34, 4.05, 1.72, title="代码层面的边界", accent=AMBER)
    add_text(slide, "layout 名称含 page-aligned，但工具中的 --alignment 当前没有真正改写 gguf ctx alignment；正确性依赖表内真实 offset/size，而不是假定 4 KiB 对齐。",
             8.77, 4.90, 3.45, 0.88, 11.2, INK, margin=0)
    finish(slide, "539f4cbf4 · tools/moe-repack/main.cpp · src/moe-offload/loader.cpp")

    # 7. Memory layout
    slide = new_slide(prs)
    add_title(slide, "共享 scratch 把“缓存容量”与“当前层可见专家轴”解耦", "BASE IMPLEMENTATION · LAYOUT", TEAL)
    add_text(slide, "6 GiB Q4 配置", 0.82, 1.58, 1.32, 0.28, 12, TEAL, True, margin=0)
    add_pill(slide, "60 persistent / layer", 2.17, 1.53, 1.72, TEAL, TEAL_LIGHT, 10)
    add_pill(slide, "196 shared scratch", 4.02, 1.53, 1.55, AMBER, AMBER_LIGHT, 10)
    add_pill(slide, "active = 256", 5.72, 1.53, 1.26, GREEN, GREEN_LIGHT, 10)

    x = 0.84
    y = 2.20
    layer_w = 1.52
    for i in range(5):
        color = TEAL if i < 4 else TEAL_LIGHT
        fill = TEAL_LIGHT if i < 4 else "E7F0F0"
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, x + i * layer_w, y, layer_w - 0.04, 0.88, fill, TEAL)
        label = f"L{i}" if i < 4 else "… L39"
        add_text(slide, label, x + i * layer_w, y + 0.15, layer_w - 0.04, 0.22, 11, color, True,
                 align=PP_ALIGN.CENTER, margin=0)
        add_text(slide, "60 slots", x + i * layer_w, y + 0.47, layer_w - 0.04, 0.20, 9.5, MUTED,
                 align=PP_ALIGN.CENTER, margin=0)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, x + 5 * layer_w, y, 3.25, 0.88, AMBER_LIGHT, AMBER)
    add_text(slide, "shared scratch · 196 slots", x + 5 * layer_w, y + 0.29, 3.25, 0.28, 12, AMBER, True,
             align=PP_ALIGN.CENTER, margin=0)
    add_text(slide, "global expert axis = 40×60 + 196 = 2,596", 0.84, 3.30, 5.22, 0.31, 13, INK, True, margin=0)
    add_text(slide, "每个 kind/dtype 各有一条同构 expert 轴", 6.32, 3.31, 4.20, 0.29, 11, MUTED, margin=0)

    add_card(slide, 0.84, 4.02, 3.64, 2.10, title="命中映射", accent=TEAL)
    add_text(slide, "global_slot = layer × 60 + local_slot", 1.10, 4.72, 3.06, 0.34, 13, TEAL, True, font=FONT_MONO, margin=0)
    add_text(slide, "跨 callback 保留；由 LRU/EAMC 管理。", 1.10, 5.27, 2.96, 0.40, 11.5, MUTED, margin=0)

    add_card(slide, 4.74, 4.02, 3.64, 2.10, title="未命中映射", accent=AMBER)
    add_text(slide, "scratch_slot = 2,400 + next", 5.00, 4.72, 3.02, 0.34, 13, AMBER, True, font=FONT_MONO, margin=0)
    add_text(slide, "只在当前 layer callback 生效；不写入持久 cache。", 5.00, 5.27, 2.96, 0.45, 11.5, MUTED, margin=0)

    add_card(slide, 8.64, 4.02, 3.86, 2.10, title="ubatch 关键含义", accent=GREEN)
    add_text(slide, "active_slots = experts = 256", 8.92, 4.70, 3.08, 0.34, 13, GREEN, True, font=FONT_MONO, margin=0)
    add_text(slide, "unique experts 不会超过 256，因此 MoE 层本身无需再按 active/top-k 限制 ubatch。", 8.92, 5.20, 3.12, 0.58, 11.3, INK, margin=0)
    finish(slide, "539f4cbf4 · docs/moe-offload-shared-scratch.md · slot_pool.cpp")

    # 8. Callback pipeline
    slide = new_slide(prs)
    add_title(slide, "每层 callback 是 offload 的控制面：先路由，再让权重就绪", "BASE IMPLEMENTATION · RUNTIME", TEAL)
    steps = [
        ("1", "GPU Top-k", "产生 [top_k, tokens] ids", BLUE),
        ("2", "D2H", "读取 ids，统计 unique / route", AMBER),
        ("3", "Cache", "hit / free / victim / scratch", TEAL),
        ("4", "I/O", "file→pinned→H2D event", RED),
        ("5", "Remap", "写 slot_ids 给 MUL_MAT_ID", GREEN),
        ("6", "Compute", "compute stream 等待后执行", VIOLET),
    ]
    for i, (num, title, body, color) in enumerate(steps):
        x = 0.73 + i * 2.08
        add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x, 1.76, 0.42, 0.42, color, color)
        add_text(slide, num, x, 1.80, 0.42, 0.23, 11, PAPER, True, align=PP_ALIGN.CENTER, margin=0)
        add_text(slide, title, x + 0.54, 1.75, 1.25, 0.28, 12.2, INK, True, margin=0)
        add_text(slide, body, x, 2.31, 1.78, 0.58, 10.4, MUTED, margin=0)
        if i < len(steps) - 1:
            add_arrow(slide, x + 1.74, 1.98, x + 2.00, 1.98, GRID, 1.3)

    add_card(slide, 0.82, 3.48, 5.65, 2.66, title="正确性不变量", accent=RED)
    add_bullets(slide, [
        "当前 call 的所有 unique expert 都必须被保护，不能互相驱逐",
        "H2D 完成事件必须进入 compute stream 的依赖链",
        "slot tensor 未使用区域先清零，避免量化 MMQ 访问垃圾数据",
        "同一层 unique 数超过 active axis 时立即 abort，不静默错算",
    ], 1.10, 4.04, 4.98, 1.75, 12.4, MUTED, RED, 6)

    add_card(slide, 6.75, 3.48, 5.74, 2.66, title="异步实现", accent=BLUE)
    add_bullets(slide, [
        "固定 pinned buffer pool + 单 worker；请求按 file offset 排序",
        "H2D 使用独立 CUDA stream；event 被 profile row 与 buffer 生命周期共同追踪",
        "提交/回收分块进行，不要求一层全部 miss 同时占用 pinned buffer",
        "请求结束补齐最后层事件并输出 CSV/summary",
    ], 7.03, 4.04, 5.07, 1.75, 12.4, MUTED, BLUE, 6)
    finish(slide, "539f4cbf4 · src/moe-offload/slot_pool.cpp · io.cpp · moe_offload_io.cu")

    # 9. Predictors
    slide = new_slide(prs)
    add_title(slide, "LRU 与 EAMC 共用准入/驱逐接口，但 EAMC 本体没有被单独证实", "PREDICTOR & PROFILER", BLUE)
    add_card(slide, 0.82, 1.56, 3.62, 4.82, title="LRU", accent=TEAL)
    add_metric(slide, "O(1)", "touch / resident lookup", 1.10, 2.14, 2.2, TEAL)
    add_bullets(slide, [
        "每层独立 slot_to_expert / exp2slot",
        "MRU 在链表头，victim 从尾部选择",
        "predictor score 平局时显式按 LRU 尾部稳定决策",
    ], 1.09, 3.34, 2.95, 1.50, 12.2, MUTED, TEAL, 7)
    add_text(slide, "正式性能矩阵与 Phase A 报告均使用 predictor=lru。", 1.09, 5.48, 2.93, 0.56, 11.2, INK, True, margin=0)

    add_card(slide, 4.70, 1.56, 3.84, 4.82, title="EAMC", accent=VIOLET)
    add_metric(slide, "EAM1", "持久化 sidecar 格式", 4.98, 2.14, 2.4, VIOLET)
    add_bullets(slide, [
        "按 request 保存稀疏 layer×expert 计数行",
        "用已观察层前缀余弦找历史近邻",
        "对目标层专家做相似度加权评分",
        "可限制有效 corpus rows；保存/加载带形状校验",
    ], 4.98, 3.34, 3.12, 1.78, 12.0, MUTED, VIOLET, 6)
    add_text(slide, "风险：评分与 sidecar 成本有 profiler 字段，但仓库内没有 EAMC vs LRU 的隔离 A/B。",
             4.98, 5.48, 3.10, 0.58, 11.1, INK, True, margin=0)

    add_card(slide, 8.80, 1.56, 3.70, 4.82, title="测量面", accent=AMBER)
    add_bullets(slide, [
        "层级：required / hit / miss / route coverage",
        "I/O：read / H2D / stall / bytes / reads",
        "控制：topk D2H / remap H2D / callback wall",
        "请求：TTFT / TPOT / total / VRAM / DRAM",
        "后续分支增加 candidate-only 与 prefill 子阶段",
    ], 9.08, 2.24, 3.04, 2.48, 12.0, MUTED, AMBER, 6)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 9.08, 5.24, 3.03, 0.76, RED_LIGHT, RED, True)
    add_text(slide, "branch 名称 ≠ 已验证变量", 9.25, 5.43, 2.70, 0.26, 12.2, RED, True,
             align=PP_ALIGN.CENTER, margin=0)
    finish(slide, "539f4cbf4 / 8b4de7c0e · predictor.cpp · profiler.cpp · Phase A report")

    # 10. Measurement hygiene
    slide = new_slide(prs)
    add_title(slide, "哪些数字可以下结论，哪些只能当线索", "MEASUREMENT HYGIENE", AMBER)
    headers = ["证据", "实验", "控制方式", "本复盘用途"]
    rows = [
        ["A", "频次准入", "old/new 二进制；同 GPU/model/page cache；重复确定", "可下因果结论"],
        ["A", "DRAM stream / P0", "同最终二进制，环境变量开关；3 次 cold repeat", "可判定退化"],
        ["A-", "ub512 vs 1024", "交替顺序 512/1024；各 4 次", "可判断自然 prompt 趋势"],
        ["B", "0714 ub / cache matrix", "单次矩阵，配置和 profiler 完整", "可看结构性趋势"],
        ["C", "MTP", "单次；Q8；无同构 no-MTP 汇总", "只报告功能与现象"],
    ]
    add_table(slide, headers, rows, 0.82, 1.66, 11.66, 3.34, [0.75, 2.35, 4.35, 3.1], highlight_rows={0, 1})
    add_card(slide, 0.82, 5.30, 5.60, 1.05, fill=RED_LIGHT, line=RED, accent=RED)
    add_text(slide, "0719 的早期误导", 1.08, 5.54, 1.50, 0.25, 11, RED, True, margin=0)
    add_text(slide, "初版 baseline 受可见 GPU / VRAM 口径干扰；最终 3-repeat baseline 把 TPOT 从 42.04 修正到约 26 ms。",
             2.65, 5.46, 3.42, 0.48, 10.7, INK, margin=0)
    add_card(slide, 6.68, 5.30, 5.80, 1.05, fill=AMBER_LIGHT, line=AMBER, accent=AMBER)
    add_text(slide, "正确性门槛", 6.95, 5.54, 1.20, 0.25, 11, AMBER, True, margin=0)
    add_text(slide, "generation hash 只证明同一模式的 repeat 一致；必须再比较 baseline 与优化模式的 token 序列。",
             8.22, 5.46, 3.86, 0.48, 10.7, INK, margin=0)
    finish(slide, "phase-a-report.md；prefill-expert-stream-*/RESULTS.md；0719 final summaries")

    # 11. 0714 ubatch result
    slide = new_slide(prs)
    add_title(slide, "0714：全活动轴释放了大 ubatch，prefill 成为首个明确收益", "0714_UBATCH_NEW", TEAL)
    add_pill(slide, "6 GiB cache · pp1024/tg1024 · Q4_K", 0.82, 1.51, 2.96, TEAL, TEAL_LIGHT, 10)
    add_bar_chart(slide, ["ub128", "ub256", "ub512", "ub1024"],
                  [("prefill ms/token", [8.47, 6.62, 5.69, 3.83], TEAL),
                   ("decode ms/token", [29.28, 29.32, 29.12, 26.90], BLUE)],
                  0.78, 2.02, 7.38, 3.58, 32, 9, "{:.1f}")
    add_card(slide, 8.46, 1.52, 4.02, 4.84, title="从 ub128 到 ub1024", accent=TEAL)
    add_metric(slide, "-54.8%", "prefill ms/token", 8.78, 2.18, 1.80, TEAL, "8.47 → 3.83")
    add_metric(slide, "-18.6%", "端到端总时长", 10.52, 2.18, 1.60, GREEN, "38.66s → 31.47s")
    add_metric(slide, "+0.82 GB", "进程峰值估算", 8.78, 3.61, 1.94, AMBER, "8.92 → 9.74 GB")
    add_metric(slide, "42.7% → 0%", "prefill unique hit", 10.52, 3.61, 1.65, RED, "一次大 batch 仍更快")
    add_text(slide, "解释", 8.78, 5.10, 0.54, 0.24, 10.5, TEAL, True, margin=0)
    add_text(slide, "共享 scratch 保证 256 专家轴完整；更少的图构建/callback 与更大的矩阵并行度，压过了 prefill cache hit 下降。",
             9.40, 5.02, 2.58, 0.73, 11.2, INK, margin=0)
    add_text(slide, "注意：这是单次矩阵（证据 B）；后续自然 prompt 交替复测确认方向。",
             0.82, 6.34, 7.32, 0.34, 11, MUTED, margin=0)
    finish(slide, "0714_ubatch_new@539f4cbf4 · combined-axis.../vram-6000/*summary.txt")

    # 12. Cache budget
    slide = new_slide(prs)
    add_title(slide, "0714：增加 2 GiB 专家缓存，decode I/O 明显下降", "CACHE BUDGET TRADE-OFF", TEAL)
    add_text(slide, "同为 ub1024 / pp1024 / tg1024", 0.82, 1.53, 3.20, 0.28, 11.5, MUTED, margin=0)
    headers = ["专家缓存", "persistent", "进程峰值", "decode TPOT", "decode hit", "SSD / token", "总时长"]
    rows = [
        ["6 GiB", "60 / layer", "9.74 GB", "26.90 ms", "92.6%", "43.08 MiB", "31.47 s"],
        ["8 GiB", "81 / layer", "11.85 GB", "22.11 ms", "95.0%", "29.26 MiB", "27.13 s"],
    ]
    add_table(slide, headers, rows, 0.82, 1.94, 11.66, 1.56,
              [1.2, 1.65, 1.55, 1.65, 1.35, 1.65, 1.45], highlight_rows={1})
    add_card(slide, 0.82, 3.94, 3.58, 2.16, title="换来的收益", accent=GREEN)
    add_metric(slide, "-17.8%", "decode TPOT", 1.10, 4.60, 1.43, GREEN, "26.90 → 22.11 ms")
    add_metric(slide, "-32.1%", "SSD bytes/token", 2.66, 4.60, 1.43, TEAL, "43.08 → 29.26 MiB")
    add_card(slide, 4.68, 3.94, 3.58, 2.16, title="付出的代价", accent=AMBER)
    add_metric(slide, "+2.11 GB", "进程峰值估算", 4.98, 4.60, 1.50, AMBER, "9.74 → 11.85 GB")
    add_metric(slide, "+14.6%", "单次 TTFT/token", 6.54, 4.60, 1.42, RED, "3.83 → 4.39 ms")
    add_card(slide, 8.54, 3.94, 3.94, 2.16, title="决策含义", accent=BLUE)
    add_text(slide, "如果目标是长输出，缓存预算优先影响 decode：2 GiB 换来 13.8% 端到端缩短。若目标是短 prompt-only，必须单独复测 TTFT。",
             8.86, 4.67, 3.28, 0.90, 12.0, INK, margin=0)
    finish(slide, "0714_ubatch_new@539f4cbf4 · vram-6000 / vram-8000 ub1024 summaries")

    # 13. Admission algorithm
    slide = new_slide(prs)
    add_title(slide, "0716：把“先看到谁”改成“谁覆盖更多路由”", "0716_EAMC · PHASE A", GREEN)
    add_pill(slide, "正式实验 predictor = LRU", 0.82, 1.50, 1.94, RED, RED_LIGHT, 10)
    add_text(slide, "因此本页验证的是 admission/tie-break，不是 EAMC 相似度预测本身。", 2.92, 1.53, 6.56, 0.28, 11, MUTED, margin=0)
    stages = [
        (0.84, "原始 top-k ids", "保留重复次数", BLUE),
        (3.11, "稳定排序", "count↓ → first↑ → id↑", GREEN),
        (5.62, "准入", "热点 miss 先占空 slot", TEAL),
        (8.11, "MRU 更新", "低频先 touch，高频最后", VIOLET),
        (10.55, "驱逐", "score 平局从 LRU 尾部", AMBER),
    ]
    for i, (x, title, body, color) in enumerate(stages):
        add_card(slide, x, 2.15, 1.78, 1.34, fill=PAPER, accent=color)
        add_text(slide, title, x + 0.18, 2.40, 1.42, 0.28, 12.4, color, True, margin=0)
        add_text(slide, body, x + 0.18, 2.85, 1.40, 0.35, 10.2, MUTED, margin=0)
        if i < len(stages) - 1:
            add_arrow(slide, x + 1.83, 2.80, stages[i + 1][0] - 0.06, 2.80, GRID, 1.4)
    add_card(slide, 0.84, 4.02, 5.56, 1.98, title="为什么只改善冷启动", accent=GREEN)
    add_text(slide, "原策略在空 cache 时按 unordered_set / 首遇顺序填满 60 个 slot，可能把低覆盖专家留在持久区。频次排序修复的是初始 cache 形状；稳态路由序列不变，因此 128 token 后几乎重合。",
             1.14, 4.66, 4.92, 0.88, 11.8, INK, margin=0)
    add_card(slide, 6.70, 4.02, 5.78, 1.98, title="新增可观测性", accent=BLUE)
    add_text(slide, "routes_required / hit / persistent，empty/victim/scratch admission，以及 route_rank_us。它们把 unique-expert miss 与实际 token route coverage 分开，避免高频专家和低频专家被等权统计。",
             7.00, 4.66, 5.12, 0.88, 11.8, INK, margin=0)
    finish(slide, "8b4de7c0e + c09907bec · admission.cpp · slot_pool.cpp · test-moe-admission.cpp")

    # 14. Admission results
    slide = new_slide(prs)
    add_title(slide, "0716：收益集中在前 128 token，整体是“小而确定”", "PHASE A RESULTS", GREEN)
    add_pill(slide, "pp1024/tg1024 · 6 GiB · ub1024 · 3 repeats", 0.82, 1.49, 3.34, GREEN, GREEN_LIGHT, 10)
    add_metric(slide, "1112 → 694", "前 8 token miss", 0.86, 2.05, 2.20, GREEN, "-37.6%")
    add_metric(slide, "24219 → 23738", "全 decode unique miss", 3.10, 2.05, 2.48, TEAL, "-1.99%")
    add_metric(slide, "43.11 → 42.22", "SSD MiB / token", 5.66, 2.05, 2.20, BLUE, "-2.06%")
    add_metric(slide, "26.73 → 26.56", "TPOT ms", 7.95, 2.05, 1.92, AMBER, "-0.64%")
    add_metric(slide, "-0.88%", "端到端时间", 10.15, 2.05, 1.75, RED, "31.47s → 31.19s")

    add_card(slide, 0.82, 3.60, 7.40, 2.42, title="decode 窗口 hit rate", accent=GREEN)
    headers = ["窗口", "旧策略", "Phase A", "变化"]
    rows = [
        ["token 0", "7.50%", "89.69%", "+82.19 pp"],
        ["1–7", "63.57%", "70.49%", "+6.92 pp"],
        ["8–31", "77.58%", "78.28%", "+0.70 pp"],
        ["32–127", "81.73%", "81.76%", "+0.03 pp"],
        ["128–1023", "94.50%", "94.50%", "0"],
    ]
    add_table(slide, headers, rows, 1.10, 4.20, 6.84, 1.54, [1.45, 1.55, 1.55, 1.45], highlight_rows={0, 1})
    add_card(slide, 8.51, 3.60, 3.97, 2.42, title="结论", accent=TEAL)
    add_text(slide, "这个方向值得合并进后续实验基线，但不值得继续围绕准入规则微调。下一阶收益必须来自 prefetch、跨层预测或更大的传输-计算重叠。",
             8.84, 4.30, 3.30, 1.12, 12.2, INK, margin=0)
    add_text(slide, "验证：41,000 profile rows 全部满足 hit+miss=required 与 coverage 边界。",
             8.84, 5.55, 3.20, 0.32, 10.4, MUTED, margin=0)
    finish(slide, "0716_eamc@5dd375005 · phase-a-report.md · early-decode-comparison.csv")

    # 15. Extended ubatch tests
    slide = new_slide(prs)
    add_title(slide, "0716：自然 prompt 与 4K prompt 复测后，ub1024 是更稳妥的整体点", "EXTENDED UBATCH TESTS", BLUE)
    add_card(slide, 0.82, 1.55, 5.68, 4.84, title="自然 prompt：ub512 vs ub1024 交替复测", accent=BLUE)
    add_bar_chart(slide, ["ub512", "ub1024"],
                  [("prefill ms/token", [5.505, 4.287], BLUE),
                   ("decode ms/token", [28.758, 27.462], TEAL)],
                  1.00, 2.23, 4.96, 2.50, 32, 9, "{:.2f}")
    add_metric(slide, "-22.1%", "prefill", 1.12, 5.10, 1.42, BLUE, "4 runs / group")
    add_metric(slide, "-4.5%", "decode mean", 2.78, 5.10, 1.50, TEAL, "ub512 variance 较大")
    add_metric(slide, "+0.49 GB", "VRAM", 4.52, 5.10, 1.42, AMBER, "9.25 → 9.74")

    add_card(slide, 6.78, 1.55, 5.70, 4.84, title="4K prompt：TTFT 最快不等于请求最快", accent=AMBER)
    headers = ["ub", "prefill ms/tok", "decode ms/tok", "峰值 GB", "估算总时长"]
    rows = [
        ["256", "5.78", "24.92", "9.06", "49.2 s"],
        ["512", "5.04", "29.91", "9.31", "51.3 s"],
        ["1024", "3.13", "24.71", "9.80", "38.1 s"],
        ["2048", "2.19", "36.42", "10.76", "46.3 s"],
        ["4096", "1.56", "32.72", "12.75", "39.9 s"],
    ]
    add_table(slide, headers, rows, 7.03, 2.15, 5.18, 2.82,
              [0.65, 1.35, 1.35, 1.0, 1.25], highlight_rows={2})
    add_text(slide, "ub4096 把 prefill 推到 1.56 ms/token，但 persistent hit=0、VRAM 与 decode 代价上升；ub1024 的端到端估算最好。",
             7.06, 5.34, 4.95, 0.65, 11.3, INK, margin=0)
    finish(slide, "4306ebc28 · 512vs1024/aggregate-summary.md；3821d9835 · input-4096 summaries")

    # 16. Performance ceiling
    slide = new_slide(prs)
    add_title(slide, "0716：全驻留性能说明 offload 的主要差距仍在数据面", "NON-OFFLOAD CEILING", RED)
    add_pill(slide, "同一 L20 / 同 Q4_K 模型；工具不同，仅作天花板", 0.82, 1.48, 3.78, RED, RED_LIGHT, 10)
    add_card(slide, 0.82, 2.08, 5.64, 3.88, title="吞吐对比", accent=RED)
    add_hbar_chart(slide, ["全驻留 pp", "offload pp", "全驻留 tg", "offload tg"],
                   [5102.37, 233.26, 158.95, 36.41], 1.06, 2.80, 4.98, 2.42, RED, 5400, " tok/s")
    add_card(slide, 6.76, 2.08, 5.72, 3.88, title="如何理解这个差距", accent=AMBER)
    add_metric(slide, "≈21.9×", "prefill 吞吐差", 7.08, 2.76, 1.74, RED, "5102 vs ≈233 tok/s")
    add_metric(slide, "≈4.4×", "decode 吞吐差", 9.12, 2.76, 1.74, AMBER, "159 vs ≈36 tok/s")
    add_bullets(slide, [
        "prefill：大量专家首次换入 + callback/graph 分段",
        "decode：miss 的 SSD read + H2D + 同步持续暴露",
        "这直接催生了后续投机与传输重叠尝试",
    ], 7.08, 4.22, 4.80, 1.30, 11.8, MUTED, RED, 7)
    add_text(slide, "不能用“模型文件 20.60 GiB”与 offload 进程峰值直接计算节省比例；全驻留结果未记录同口径进程峰值。",
             0.82, 6.31, 11.52, 0.36, 11, MUTED, margin=0)
    finish(slide, "5dd375005 · non-offload-single-gpu/llama-bench-pp1024-tg1024.md；512vs1024 retest")

    # 17. Ngram speculative architecture
    slide = new_slide(prs)
    add_title(slide, "0719：n-gram 投机减少 target step，却放大每步专家集合", "SPECULATIVE DECODING · DESIGN", RED)
    add_card(slide, 0.82, 1.56, 5.66, 2.06, title="验证循环", accent=RED)
    add_text(slide, "committed token", 1.12, 2.22, 1.52, 0.30, 12, TEAL, True, margin=0)
    add_text(slide, "+", 2.67, 2.20, 0.24, 0.30, 14, MUTED, True, margin=0)
    add_text(slide, "n-gram candidates", 2.95, 2.22, 1.65, 0.30, 12, RED, True, margin=0)
    add_arrow(slide, 4.72, 2.38, 5.42, 2.38, MUTED, 1.6)
    add_text(slide, "一次 target batch", 5.02, 2.72, 1.06, 0.28, 10.5, INK, True, align=PP_ALIGN.CENTER, margin=0)
    add_text(slide, "接受前缀；回滚未接受位置", 1.12, 2.85, 3.40, 0.32, 11, MUTED, margin=0)

    add_card(slide, 6.78, 1.56, 5.70, 2.06, title="offload 视角", accent=AMBER)
    add_text(slide, "target step 数下降", 7.08, 2.22, 1.55, 0.30, 12, GREEN, True, margin=0)
    add_text(slide, "vs", 8.80, 2.22, 0.35, 0.30, 11, MUTED, True, margin=0)
    add_text(slide, "candidate-only expert 增加", 9.24, 2.22, 2.20, 0.30, 12, RED, True, margin=0)
    add_text(slide, "收益条件：减少的 target 启动/计算 > 额外 SSD/H2D + cache 污染 + rollback。",
             7.08, 2.83, 4.74, 0.42, 11.3, INK, margin=0)

    headers = ["阶段", "处理方式", "目标"]
    rows = [
        ["Stage 1/2", "所有 candidate 参与正常准入；新增 candidate-only 指标", "先测额外 I/O"],
        ["Stage 3", "candidate-only miss 只进 scratch；predictor 仅观察 committed", "避免 cache 污染"],
        ["Gate", "各 16 outputs 校准；spec < 0.98×base 才继续", "运行时自适应关闭"],
    ]
    add_table(slide, headers, rows, 0.82, 4.18, 8.15, 1.72, [1.25, 4.55, 2.05], highlight_rows={1})
    add_card(slide, 9.25, 4.18, 3.23, 1.72, title="配套修复", accent=VIOLET)
    add_text(slide, "• sequence rollback capability probe\n• partial checkpoint replay\n• per-mode generation hash\n• candidate-only bytes/hit 指标",
             9.53, 4.74, 2.63, 0.92, 10.8, MUTED, margin=0)
    finish(slide, "0719_speculative_decoding@94fdf126b · tools/moe-bench/main.cpp · slot_pool.cpp")

    # 18. Ngram results
    slide = new_slide(prs)
    add_title(slide, "0719：最终重复数据中，投机解码更慢、读得更多，而且 token 等价性未通过", "SPECULATIVE DECODING · RESULTS", RED)
    add_pill(slide, "final 3-repeat · pp1024/tg1024 · 6 GiB", 0.82, 1.48, 2.83, RED, RED_LIGHT, 10)
    headers = ["模式", "TPOT", "相对基线", "SSD MiB/tok", "accept", "verify step", "状态"]
    rows = [
        ["Baseline", "26.40 ms", "—", "42.22", "—", "1.00 out/step", "基线"],
        ["Stage 2", "31.20 ms", "+18.2%", "61.29", "63.96%", "1.43 out/step", "退化"],
        ["Stage 3 gate", "34.11 ms", "+29.2%", "65.45", "0%*", "1.00 out/step", "自动关闭"],
    ]
    add_table(slide, headers, rows, 0.82, 1.96, 11.66, 1.86,
              [1.55, 1.25, 1.30, 1.50, 1.12, 1.65, 1.40], highlight_rows={1, 2})
    add_card(slide, 0.82, 4.15, 3.72, 1.88, title="额外数据面成本", accent=RED)
    add_metric(slide, "+45.2%", "Stage 2 SSD/token", 1.12, 4.83, 1.52, RED, "42.22 → 61.29 MiB")
    add_metric(slide, "20.57 MiB", "candidate-only / output", 2.83, 4.83, 1.38, AMBER, "Stage 2")
    add_card(slide, 4.82, 4.15, 3.72, 1.88, title="Stage 3 的判定", accent=AMBER)
    add_metric(slide, "94.61 ms", "spec 校准 / output", 5.12, 4.83, 1.40, RED)
    add_metric(slide, "45.40 ms", "base 校准 / output", 6.82, 4.83, 1.40, GREEN, "3/3 repeats 关闭")
    add_card(slide, 8.82, 4.15, 3.66, 1.88, title="正确性红线", accent=RED)
    add_text(slide, "final baseline / stage2 / stage3 的 generation hash 不同；token CSV 从约第 101 个输出开始分叉。repeats_consistent=true 只代表各模式内部稳定。",
             9.12, 4.76, 3.05, 0.92, 11.2, INK, True, margin=0)
    add_text(slide, "* Stage 3 仅在校准期尝试 48 个 draft，全部未接受，之后回到 baseline 路径。",
             0.84, 6.38, 11.40, 0.29, 10.6, MUTED, margin=0)
    finish(slide, "94fdf126b · 0720-speculative-stage3/final-*.summary.txt + final-*-tokens.csv")

    # 19. MTP design
    slide = new_slide(prs)
    add_title(slide, "0720：MTP 走模型内 nextn 路径，并补上 CUDA MMQ 的尾块边界", "MTP OFFLOAD · IMPLEMENTATION", VIOLET)
    add_card(slide, 0.82, 1.55, 5.55, 4.92, title="双 context 推理路径", accent=VIOLET)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 1.13, 2.24, 1.72, 0.74, BLUE_LIGHT, BLUE, True)
    add_text(slide, "target ctx", 1.13, 2.45, 1.72, 0.27, 13, BLUE, True, align=PP_ALIGN.CENTER, margin=0)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, 3.92, 2.24, 1.72, 0.74, VIOLET_LIGHT, VIOLET, True)
    add_text(slide, "MTP draft ctx", 3.92, 2.45, 1.72, 0.27, 13, VIOLET, True, align=PP_ALIGN.CENTER, margin=0)
    add_arrow(slide, 2.88, 2.61, 3.84, 2.61, MUTED, 1.7)
    add_text(slide, "ctx_other", 3.11, 2.29, 0.62, 0.24, 9, MUTED, margin=0)
    add_bullets(slide, [
        "target prefill 后 common_speculative_process 初始化 MTP",
        "draft_n_max=1 为主；target 一次验证 token + draft",
        "accept 后同步 target/draft 的 KV tail；n_rs_seq 至少 1",
        "Q8 KV cache 与独立 draft perf 计数",
    ], 1.12, 3.42, 4.82, 2.02, 11.8, MUTED, VIOLET, 6)
    add_text(slide, "MTP repack 记录 41 个 MoE 层；6 GiB 预算对应 45 persistent slots/layer。",
             1.12, 5.76, 4.78, 0.41, 11.1, INK, True, margin=0)

    add_card(slide, 6.68, 1.55, 5.80, 2.22, title="MMQ 尾块修复", accent=RED)
    add_text(slide, "问题", 6.98, 2.19, 0.45, 0.24, 10.5, RED, True, margin=0)
    add_text(slide, "多输出 candidate batch 的 flattened row 数不一定是 J_max 的倍数，旧 kernel 会越界读 ids_dst。",
             7.52, 2.11, 4.38, 0.56, 11.4, INK, margin=0)
    add_text(slide, "修复", 6.98, 2.91, 0.45, 0.24, 10.5, GREEN, True, margin=0)
    add_text(slide, "把 ids 与量化 workspace pad 到 J_max；kernel 对 col < col_diff 做保护。",
             7.52, 2.83, 4.36, 0.50, 11.4, INK, margin=0)

    add_card(slide, 6.68, 4.05, 5.80, 2.42, title="指标修复", accent=BLUE)
    add_text(slide, "原 profiler 用 n_tokens>1 推断 prefill，会把 MTP process / speculative decode 混入 prefill。最新提交改为显式 request phase：",
             6.98, 4.65, 5.12, 0.63, 11.4, INK, margin=0)
    add_pill(slide, "prefill_target", 6.99, 5.48, 1.35, BLUE, BLUE_LIGHT, 9)
    add_pill(slide, "prefill_mtp_process", 8.52, 5.48, 1.72, VIOLET, VIOLET_LIGHT, 9)
    add_pill(slide, "decode", 10.42, 5.48, 0.88, TEAL, TEAL_LIGHT, 9)
    add_text(slide, "summary 同时保留 prefill_total，避免修复前后口径断裂。", 6.99, 5.97, 4.88, 0.28, 10.6, MUTED, margin=0)
    finish(slide, "0edf6febc + fde549896 · tools/moe-bench/main.cpp · ggml-cuda/mmq.cu/.cuh")

    # 20. MTP results
    slide = new_slide(prs)
    add_title(slide, "0720：MTP 已跑通，但当前数字不能证明相对 no-MTP 的净收益", "MTP OFFLOAD · RESULTS", VIOLET)
    add_pill(slide, "Q8_0 · 6 GiB · 45/256 slots · 单次", 0.82, 1.49, 2.84, VIOLET, VIOLET_LIGHT, 10)
    headers = ["结果集", "TTFT", "TPOT", "accept", "decode hit", "SSD/tok", "说明"]
    rows = [
        ["首轮正式", "6.58 s", "8.99 ms", "83.3%", "98.7%", "0.20 MiB", "旧 phase 口径"],
        ["normal request", "8.69 s", "9.46 ms", "87.0%", "96.2%", "0.55 MiB", "prompt 不同"],
        ["prefill split", "11.38 s", "9.39 ms", "94.9%", "92.4%", "3.44 MiB", "新 phase 口径"],
    ]
    add_table(slide, headers, rows, 0.82, 1.95, 11.66, 2.04,
              [1.55, 1.15, 1.15, 1.10, 1.20, 1.30, 2.20], highlight_rows={2})
    add_card(slide, 0.82, 4.33, 3.72, 1.76, title="最新 TTFT 分解", accent=BLUE)
    add_metric(slide, "11.074 s", "target prefill", 1.12, 4.93, 1.46, BLUE)
    add_metric(slide, "0.306 s", "MTP process", 2.84, 4.93, 1.34, VIOLET, "只占 2.7%")
    add_card(slide, 4.82, 4.33, 3.72, 1.76, title="积极信号", accent=GREEN)
    add_text(slide, "MTP draft/accept、KV rollback、offload callback 和 MMQ 多输出路径已经在同一工具中闭环；decode I/O 在部分 run 中很低。",
             5.13, 4.95, 3.06, 0.78, 11.5, INK, margin=0)
    add_card(slide, 8.82, 4.33, 3.66, 1.76, title="证据缺口", accent=RED)
    add_text(slide, "无同 commit、同 Q8、同 prompt 的 no-MTP full summary；单次波动大；仓库中也没有 baseline/MTP token 序列等价报告。",
             9.12, 4.95, 3.02, 0.78, 11.5, INK, True, margin=0)
    add_text(slide, "因此本页只说明“功能跑通 + 指标可解释”，不把 8.99 ms/token 写成已证实加速。",
             0.83, 6.39, 11.43, 0.29, 10.8, MUTED, margin=0)
    finish(slide, "0720_mtp@fde549896 · mtp-pp1024-tg1024/**/summary.txt")

    # 21. Transfer phase1 design
    slide = new_slide(prs)
    add_title(slide, "0724 Phase 1：先让 gate weight 边到边算，up/down 仍整批等待", "DRAM STREAMING · PHASE 1", AMBER)
    add_pill(slide, "仅 prefill + DRAM source + tokens>8 + miss>1", 0.82, 1.49, 3.50, AMBER, AMBER_LIGHT, 10)
    steps = [
        (0.84, "DRAM mmap", "只读映射 + 预触页", BLUE),
        (3.14, "Callback", "按 kind / route rank 提交", TEAL),
        (5.64, "H2D stream", "每 (slot,kind) 发布 event", RED),
        (8.14, "Gate MMQ", "ready group≤8 分组执行", AMBER),
        (10.64, "Up / Down", "等待完整 kind，保持 batched", VIOLET),
    ]
    for i, (x, title, body, color) in enumerate(steps):
        add_card(slide, x, 2.17, 1.80, 1.38, fill=PAPER, accent=color)
        add_text(slide, title, x + 0.18, 2.43, 1.44, 0.28, 12.2, color, True, margin=0)
        add_text(slide, body, x + 0.18, 2.88, 1.42, 0.38, 10.0, MUTED, margin=0)
        if i < len(steps) - 1:
            add_arrow(slide, x + 1.85, 2.85, steps[i + 1][0] - 0.06, 2.85, GRID, 1.4)
    add_card(slide, 0.84, 4.10, 5.64, 1.94, title="希望隐藏的成本", accent=GREEN)
    add_text(slide, "等待一整层 3×miss blobs", 1.15, 4.81, 2.15, 0.32, 13, RED, True, margin=0)
    add_arrow(slide, 3.36, 4.98, 4.14, 4.98, MUTED, 1.7)
    add_text(slide, "gate miss ready 即算", 4.27, 4.81, 1.62, 0.32, 13, GREEN, True, margin=0)
    add_text(slide, "理论上用 gate 计算覆盖后续 up/down 传输。", 1.15, 5.35, 4.56, 0.32, 11, MUTED, margin=0)
    add_card(slide, 6.78, 4.10, 5.70, 1.94, title="Phase 1 仍保留的控制开销", accent=RED)
    add_bullets(slide, [
        "expert_bounds device→host copy + cudaStreamSynchronize",
        "host sort + 多个小 gate MMQ launch + 每组等待",
        "只流水 gate；完整 expert 依赖链没有被 overlap",
    ], 7.08, 4.72, 5.04, 1.08, 11.6, MUTED, RED, 5)
    finish(slide, "0724_transfer@2dcac38ea · mmq.cu · moe_offload_io.cu · io.cpp · slot_pool.cpp")

    # 22. Phase1 results
    slide = new_slide(prs)
    add_title(slide, "0724 Phase 1：控制面等待少了，但 GPU 时间线更差", "DRAM STREAMING · PHASE 1 RESULTS", AMBER)
    add_pill(slide, "同最终二进制开关 · DRAM prefault · cold cache · 3 repeats", 0.82, 1.49, 4.14, AMBER, AMBER_LIGHT, 10)
    add_bar_chart(slide, ["64", "128", "256", "512", "1024"],
                  [("batched TTFT ms", [1360.2, 1733.9, 2220.3, 2733.4, 3351.8], BLUE),
                   ("stream TTFT ms", [1613.5, 2055.2, 2650.5, 3164.9, 3695.5], RED)],
                  0.78, 2.02, 7.62, 3.72, 4000, 8, "{:.0f}")
    add_card(slide, 8.68, 1.55, 3.80, 4.85, title="所有点都未交叉", accent=RED)
    add_metric(slide, "+18.6%", "64 token", 8.98, 2.20, 1.38, RED)
    add_metric(slide, "+19.4%", "256 token", 10.58, 2.20, 1.38, RED)
    add_metric(slide, "+10.3%", "1024 token", 8.98, 3.54, 1.38, AMBER, "惩罚虽收窄")
    add_metric(slide, "0.10 ms/tok", "1024 callback wall", 10.58, 3.54, 1.50, GREEN, "2.64 → 0.10")
    add_text(slide, "但 1024 的 CUDA-timed interval 从 0.19 升到 2.95 ms/token；其中包含 stream wait 和额外小 kernel launch。",
             8.98, 4.96, 3.02, 0.78, 11.4, INK, margin=0)
    add_text(slide, "结论：被隐藏的是 host callback 阻塞，暴露出来的是串行 transfer wait 与 gate 分组启动成本。",
             0.83, 6.34, 11.42, 0.36, 11.2, MUTED, margin=0)
    finish(slide, "2dcac38ea · prefill-expert-stream-phase1/RESULTS.md + final-ab/summary.csv")

    # 23. P0 cleanup
    slide = new_slide(prs)
    add_title(slide, "0724 P0：移除 D2H、整流同步和重复排序，结论仍未改变", "TRANSFER TO CPU · P0", AMBER)
    add_card(slide, 0.82, 1.55, 5.64, 2.42, title="P0 做了什么", accent=TEAL)
    add_bullets(slide, [
        "callback 直接发布 [slot, routed_count, miss] execution plan",
        "Gate MMQ 不再读回 expert_bounds，不再 cudaStreamSynchronize",
        "miss blobs 用三类线性 bucket 替代 comparison sort",
        "无 debug 开关时不采 host timing，避免测量干扰",
    ], 1.12, 2.16, 4.95, 1.52, 11.6, MUTED, TEAL, 5)
    add_card(slide, 6.76, 1.55, 5.72, 2.42, title="诊断证明 plan 不是主要问题", accent=BLUE)
    add_metric(slide, "40", "plans / layers", 7.06, 2.16, 1.18, BLUE)
    add_metric(slide, "4,633", "active entries", 8.42, 2.16, 1.26, TEAL)
    add_metric(slide, "7.77", "mean group", 9.92, 2.16, 1.20, GREEN, "limit=8")
    add_metric(slide, "3", "singleton groups", 11.16, 2.16, 0.78, AMBER)
    add_text(slide, "128-token debug 中 host wait weight-ready 仍累计 443.7 ms。", 7.06, 3.43, 4.72, 0.30, 10.7, MUTED, margin=0)

    headers = ["tokens", "batched", "Phase 1", "P0", "P0 delta"]
    rows = [
        ["512", "2682.2 ms", "3164.9 ms*", "3131.3 ms", "+16.74%"],
        ["1024", "3322.5 ms", "3695.5 ms*", "3715.5 ms", "+11.83%"],
    ]
    add_table(slide, headers, rows, 0.82, 4.32, 7.20, 1.40, [0.95, 1.55, 1.55, 1.45, 1.40], highlight_rows={0, 1})
    add_card(slide, 8.32, 4.32, 4.16, 1.40, fill=RED_LIGHT, line=RED, accent=RED)
    add_text(slide, "停止条件已满足", 8.65, 4.62, 1.28, 0.26, 11, RED, True, margin=0)
    add_text(slide, "仍远高于 3% regression threshold；不再继续 gate-only 微优化。",
             9.98, 4.52, 2.10, 0.50, 10.8, INK, True, margin=0)
    add_text(slide, "* Phase 1 与 P0 是不同正式批次，数值用于确认结论一致，不作逐毫秒因果比较。",
             0.83, 6.33, 11.42, 0.32, 10.5, MUTED, margin=0)
    finish(slide, "a03c4fa82 · prefill-expert-stream-p0/RESULTS.md + final-ab/summary.csv")

    # 24. Synthesis matrix
    slide = new_slide(prs)
    add_title(slide, "跨分支综合：哪些进入新基线，哪些保留为失败知识", "SYNTHESIS", CHARCOAL)
    headers = ["尝试", "核心假设", "最好证据", "判断", "下一动作"]
    rows = [
        ["共享 scratch", "活动轴覆盖全部专家", "ub1024 稳定运行", "保留", "作为共同基线"],
        ["大 ubatch", "减少 callback、增大 GEMM", "TTFT/token -54.8%", "保留", "默认 ub1024"],
        ["8 GiB cache", "hit 换 I/O", "TPOT -17.8%", "按场景", "做预算曲线"],
        ["频次准入", "热点先占空 slot", "前 8 token miss -37.6%", "保留", "停止微调"],
        ["EAMC predictor", "历史相似请求预测", "无隔离 A/B", "未知", "先做 LRU 对照"],
        ["n-gram spec", "少 target step", "TPOT +18.2%；token 分叉", "否", "先修正确性"],
        ["MTP", "高质量 draft 抵消验证", "83–95% accept；无 A/B", "待验证", "同构 A/B"],
        ["Gate stream/P0", "H2D 与 gate 重叠", "TTFT +10–19%", "否", "完整 executor"],
    ]
    add_table(slide, headers, rows, 0.72, 1.53, 11.90, 4.94,
              [1.65, 2.75, 2.55, 1.10, 2.35], highlight_rows={0, 1, 3})
    # Overlay judgement pills for stronger scanability.
    judgement = [("保留", GREEN), ("保留", GREEN), ("按场景", BLUE), ("保留", GREEN),
                 ("未知", AMBER), ("否", RED), ("待验证", VIOLET), ("否", RED)]
    # The table already carries text; colored labels are summarized below to keep it editable.
    add_pill(slide, "保留", 0.82, 6.60, 0.74, GREEN, GREEN_LIGHT, 9)
    add_pill(slide, "待验证", 1.72, 6.60, 0.88, VIOLET, VIOLET_LIGHT, 9)
    add_pill(slide, "停止", 2.76, 6.60, 0.74, RED, RED_LIGHT, 9)
    add_text(slide, "最清晰的规律：凡是减少 miss/callback 数量的尝试有效；只改变等待位置、但不减少总传输/launch 的尝试无效。",
             3.72, 6.59, 8.54, 0.32, 11.2, INK, True, margin=0)
    finish(slide, "综合所有实验 summary；判断按证据等级而非 branch 名称")

    # 25. Bottleneck evolution
    slide = new_slide(prs)
    add_title(slide, "瓶颈如何在尝试中迁移", "WHAT WE LEARNED", BLUE)
    cols = [
        (0.82, "基础 offload", RED, ["大量 cold miss", "callback wall", "SSD + H2D"]),
        (3.34, "ub / admission", GREEN, ["callback 次数↓", "早期 miss↓", "稳态不变"]),
        (5.86, "投机", VIOLET, ["target step↓", "candidate I/O↑", "正确性风险"]),
        (8.38, "gate stream", AMBER, ["host wait↓", "compute wait↑", "小 launch↑"]),
        (10.90, "下一目标", TEAL, ["总 bytes↓", "ready-set batched", "完整 expert pipeline"]),
    ]
    for i, (x, title, color, bullets) in enumerate(cols):
        add_card(slide, x, 1.85, 1.88, 3.78, fill=PAPER, accent=color)
        add_text(slide, title, x + 0.20, 2.14, 1.48, 0.32, 13.5, color, True, margin=0)
        for j, item in enumerate(bullets):
            add_shape(slide, MSO_AUTO_SHAPE_TYPE.OVAL, x + 0.22, 2.92 + j * 0.72, 0.16, 0.16, color, color)
            add_text(slide, item, x + 0.48, 2.82 + j * 0.72, 1.16, 0.38, 11.0, INK, True, margin=0)
        if i < len(cols) - 1:
            add_arrow(slide, x + 1.92, 3.72, cols[i + 1][0] - 0.06, 3.72, MUTED, 1.7)
    add_card(slide, 0.82, 5.98, 11.96, 0.62, fill=BLUE_LIGHT, line=BLUE, accent=BLUE)
    add_text(slide, "性能问题没有被“异步”自动消失：只有减少必做工作，或让足够大的计算块覆盖传输，才会转化为端到端收益。",
             1.12, 6.13, 11.28, 0.26, 12.2, INK, True, align=PP_ALIGN.CENTER, margin=0)
    finish(slide, "由 0714→0716→0719/0720→0724 的结果归纳")

    # 26. Recommendations
    slide = new_slide(prs)
    add_title(slide, "下一轮建议：先把比较变可靠，再把流水做完整", "RECOMMENDED NEXT STEPS", GREEN)
    next_steps = [
        ("P0", "正确性门槛", RED, "每个优化模式必须与 greedy baseline 做 token-by-token 比较；先定位 0719 在输出约 #101 的分叉。"),
        ("P1", "统一 A/B harness", BLUE, "同 commit、同 quant、同 prompt、固定 visible GPU 与 cache reset；保存命令、环境、token hash 和显存 baseline。"),
        ("P2", "MTP 受控实验", VIOLET, "Q8 MTP on/off 同二进制；报告 target eval tokens、candidate-only bytes、acceptance 与端到端 TPOT。"),
        ("P3", "完整 expert executor", TEAL, "以 ready set 批处理 gate/up→activation→down；避免每专家小 kernel；再验证是否真正覆盖 H2D。"),
        ("P4", "减少总数据量", AMBER, "评估 host/pinned 热缓存、跨层预取、coalesced read；只有命中率与 bytes 同时改善才保留 predictor。"),
    ]
    for i, (prio, title, color, body) in enumerate(next_steps):
        y = 1.54 + i * 1.00
        add_pill(slide, prio, 0.82, y + 0.08, 0.56, color, color + "" if False else (RED_LIGHT if color == RED else BLUE_LIGHT if color == BLUE else VIOLET_LIGHT if color == VIOLET else TEAL_LIGHT if color == TEAL else AMBER_LIGHT), 10)
        add_text(slide, title, 1.60, y + 0.06, 2.05, 0.31, 13.5, color, True, margin=0)
        add_text(slide, body, 3.70, y, 8.52, 0.58, 11.6, INK, margin=0)
        if i < len(next_steps) - 1:
            line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1.08), Inches(y + 0.48),
                                              Inches(1.08), Inches(y + 1.07))
            line.line.color.rgb = rgb(GRID)
            line.line.width = Pt(1.3)
    add_card(slide, 0.82, 6.62, 11.66, 0.48, fill=GREEN_LIGHT, line=GREEN, accent=GREEN)
    add_text(slide, "成功标准：correct tokens + lower wall time + lower/controlled bytes；三者缺一不可。",
             1.16, 6.72, 10.94, 0.22, 11.8, GREEN, True, align=PP_ALIGN.CENTER, margin=0)
    finish(slide, "建议由现有瓶颈与证据缺口推导；不包含新的代码实现")

    # 27. Reproducibility/source map
    slide = new_slide(prs)
    add_title(slide, "关键文件与复现入口", "SOURCE MAP", CHARCOAL)
    headers = ["主题", "代码 / 报告", "分支 / 提交"]
    rows = [
        ["共享框架", "docs/moe-offload-shared-scratch.md；src/moe-offload/{slot_pool,io,predictor}.*", "0714 @ 539f4cbf4"],
        ["repack", "tools/moe-repack/main.cpp；models/*.moe.gguf.json", "0714 @ 539f4cbf4"],
        ["准入", "admission.*；phase-a-report.md；early-decode-comparison.csv", "0716 @ 8b4de7c0e / c09907bec"],
        ["ubatch / 上限", "new-result/**；non-offload-single-gpu/**", "0716 @ 3821d9835 / 4306ebc28 / 5dd375005"],
        ["n-gram spec", "0719-speculative-stage12/**；0720-speculative-stage3/**", "0719 @ 94fdf126b"],
        ["MTP", "mtp-pp1024-tg1024/**；ggml-cuda/mmq.*；profiler.*", "0720 @ 0edf6febc / fde549896"],
        ["DRAM stream", "prefill-expert-stream-phase1/RESULTS.md", "0724 @ 2dcac38ea"],
        ["P0", "prefill-expert-stream-p0/RESULTS.md", "0724 @ a03c4fa82"],
    ]
    add_table(slide, headers, rows, 0.72, 1.52, 11.90, 4.74, [1.55, 7.10, 3.25], highlight_rows=None)
    add_text(slide, "典型构建", 0.82, 6.50, 0.80, 0.24, 10.5, TEAL, True, margin=0)
    add_text(slide, "cmake -S . -B build-moe -DGGML_CUDA=ON -DLLAMA_MOE_OFFLOAD=ON  ·  cmake --build build-moe --target llama-moe-bench llama-moe-repack",
             1.72, 6.45, 10.45, 0.35, 9.6, INK, font=FONT_MONO, margin=0)
    finish(slide, "PPT 生成脚本：ppt/build_offload_review.py · 数据截至 2026-07-27")

    # 28. Closing
    slide = new_slide(prs, CHARCOAL)
    add_text(slide, "最终判断", 0.82, 0.72, 2.10, 0.34, 12, AMBER, True, margin=0)
    add_text(slide, "把 offload 做快的关键，不是更多异步，\n而是更少的数据和更大的可覆盖计算块。",
             0.82, 1.38, 10.90, 1.48, 27, PAPER, True, margin=0)
    add_shape(slide, MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0.82, 3.32, 11.56, 0.03, TEAL, TEAL)
    final_items = [
        ("保留", "shared scratch / ub1024 / frequency admission", GREEN),
        ("验证", "MTP on/off controlled A/B", VIOLET),
        ("停止", "n-gram 当前实现 / gate-only streaming", RED),
        ("下一步", "correctness gate + full expert executor", AMBER),
    ]
    for i, (tag, body, color) in enumerate(final_items):
        x = 0.84 + (i % 2) * 5.90
        y = 3.86 + (i // 2) * 1.20
        add_pill(slide, tag, x, y, 0.82, color,
                 GREEN_LIGHT if color == GREEN else VIOLET_LIGHT if color == VIOLET else RED_LIGHT if color == RED else AMBER_LIGHT, 10)
        add_text(slide, body, x + 1.05, y + 0.02, 4.55, 0.34, 13, PAPER, True, margin=0)
    add_text(slide, "llama.cpp.clean · MoE offload experiment review", 0.84, 6.74, 7.20, 0.24, 9.8, "AAB7BD", margin=0)
    finish(slide, "End")

    prs.save(OUT_FILE)
    return OUT_FILE, len(prs.slides)


if __name__ == "__main__":
    path, count = make_deck()
    print(f"wrote {path} ({count} slides)")
